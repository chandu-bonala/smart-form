from __future__ import annotations

import csv
import json
from datetime import datetime
from io import BytesIO, StringIO
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.answer import Answer
from app.models.form import Form
from app.models.question import Question
from app.models.response import Response


class ExportService:
    """Exports form responses to CSV, JSON, Excel, and Word-friendly analytics exports."""

    @staticmethod
    async def export_to_csv(db: AsyncSession, form_id: str) -> str:
        """Export all responses as CSV with headers."""
        form = await db.execute(
            select(Form)
            .options(selectinload(Form.questions).selectinload(Question.options))
            .where(Form.id == form_id)
        )
        form = form.scalar_one_or_none()
        if not form:
            return ""

        responses = await db.execute(
            select(Response)
            .options(selectinload(Response.answers).selectinload(Answer.question))
            .where(Response.form_id == form_id)
            .order_by(Response.submitted_at.desc())
        )
        responses = list(responses.scalars().unique().all())

        # Build CSV with a UTF-8 BOM so Excel opens it correctly on Windows.
        output = StringIO(newline="")
        output.write("\ufeff")
        
        # Create header row
        headers = ["response_id", "submitted_at", "completion_time (seconds)", "device", "browser", "ip_address"]
        for q in form.questions:
            headers.append(q.title)

        writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()

        # Write response rows
        for response in responses:
            row = {
                "response_id": response.id,
                "submitted_at": response.submitted_at.isoformat(),
                "completion_time (seconds)": response.completion_time or "",
                "device": response.device or "",
                "browser": response.browser or "",
                "ip_address": response.ip_address or "",
            }

            for answer in response.answers:
                question_title = answer.question.title if answer.question else (answer.question_snapshot or {}).get("title")
                if question_title:
                    row[question_title] = answer.answer or ""

            writer.writerow(row)

        return output.getvalue()

    @staticmethod
    async def export_to_json(db: AsyncSession, form_id: str) -> str:
        """Export all responses as JSON."""
        form = await db.execute(
            select(Form)
            .options(selectinload(Form.questions))
            .where(Form.id == form_id)
        )
        form = form.scalar_one_or_none()
        if not form:
            return "{}"

        responses = await db.execute(
            select(Response)
            .options(selectinload(Response.answers).selectinload(Answer.question))
            .where(Response.form_id == form_id)
            .order_by(Response.submitted_at.desc())
        )
        responses = list(responses.scalars().unique().all())

        export_data = {
            "form_id": form.id,
            "form_title": form.title,
            "total_responses": len(responses),
            "exported_at": datetime.now().isoformat(),
            "responses": []
        }

        for response in responses:
            response_data = {
                "response_id": response.id,
                "submitted_at": response.submitted_at.isoformat(),
                "completion_time_seconds": response.completion_time,
                "device": response.device,
                "browser": response.browser,
                "ip_address": response.ip_address,
                "answers": {}
            }

            for answer in response.answers:
                question_title = answer.question.title if answer.question else (answer.question_snapshot or {}).get("title")
                if question_title:
                    response_data["answers"][question_title] = answer.answer

            export_data["responses"].append(response_data)

        return json.dumps(export_data, indent=2, default=str)

    @staticmethod
    async def export_to_xlsx(db: AsyncSession, form_id: str) -> bytes:
        """Export all responses as an Excel workbook."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise RuntimeError('The openpyxl package is required for Excel export.') from exc

        form = await db.execute(
            select(Form)
            .options(selectinload(Form.questions).selectinload(Question.options))
            .where(Form.id == form_id)
        )
        form = form.scalar_one_or_none()
        if not form:
            return b""

        responses = await db.execute(
            select(Response)
            .options(selectinload(Response.answers).selectinload(Answer.question))
            .where(Response.form_id == form_id)
            .order_by(Response.submitted_at.desc())
        )
        responses = list(responses.scalars().unique().all())

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Responses'
        sheet.freeze_panes = "A2"
        headers = ["response_id", "submitted_at", "completion_time_seconds", "device", "browser", "ip_address"]
        headers.extend(question.title for question in form.questions)
        sheet.append(headers)
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

        header_fill = PatternFill("solid", fgColor="5B4BFF")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for response in responses:
            row = {
                "response_id": response.id,
                "submitted_at": response.submitted_at.isoformat(),
                "completion_time_seconds": response.completion_time or "",
                "device": response.device or "",
                "browser": response.browser or "",
                "ip_address": response.ip_address or "",
            }
            for answer in response.answers:
                question_title = answer.question.title if answer.question else (answer.question_snapshot or {}).get("title")
                if question_title:
                    row[question_title] = answer.answer or ""
            sheet.append([row.get(header, "") for header in headers])

        # Make the sheet easier to read in Excel.
        widths = {
            "A": 40,
            "B": 24,
            "C": 24,
            "D": 16,
            "E": 16,
            "F": 18,
        }
        for column_letter, width in widths.items():
            sheet.column_dimensions[column_letter].width = width
        for column_index in range(7, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(column_index)].width = 24
        for column_cells in sheet.iter_cols(min_row=2):
            for cell in column_cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def export_to_docx(form: Form, analytics_data: Any) -> bytes:
        """Export a clean Word document summary of form analytics."""
        try:
            from docx import Document
            from docx.shared import Inches, Pt
        except ImportError as exc:
            raise RuntimeError('The python-docx package is required for Word export.') from exc

        document = Document()
        document.add_heading(form.title, level=0)
        document.add_paragraph(f"Form ID: {form.id}")

        summary = document.add_table(rows=4, cols=2)
        summary.style = "Table Grid"
        summary_rows = [
            ("Total responses", analytics_data.total_responses),
            ("Completion rate", f"{analytics_data.completion_rate}%"),
            ("Average completion time", f"{analytics_data.average_completion_time} seconds"),
            ("Questions", len(getattr(analytics_data, "questions", []))),
        ]
        for row_index, (label, value) in enumerate(summary_rows):
            summary.cell(row_index, 0).text = str(label)
            summary.cell(row_index, 1).text = str(value)

        for row in summary.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10.5)
            row.cells[0].width = Inches(2.3)
            row.cells[1].width = Inches(4.8)

        document.add_heading('Question breakdown', level=1)
        for question in analytics_data.questions:
            document.add_paragraph(f"{question.title}: {question.total_answers} answers")
            if question.options:
                for option in question.options:
                    document.add_paragraph(f"{option.option}: {option.count} ({option.percentage}%)", style="List Bullet")
        output = BytesIO()
        document.save(output)
        return output.getvalue()
